from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db.models import Avg
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_POST
from .models import Article, ArticleRating, Comment, Bookmark, Notification, Subscription
from .forms import ArticleForm, SignUpForm, CommentForm
from django.contrib.auth import login
from django.contrib.auth.models import User
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def article_list(request):
    articles = Article.objects.filter(is_published=True)
    
    return render(
        request,
        "articles/article_list.html",
        {"articles" : articles}
    )


def article_detail(request, pk):
    article = get_object_or_404(Article, pk=pk)
    ratings = ArticleRating.objects.filter(article=article)

    # --- РАСЧЕТ ВРЕМЕНИ ЧТЕНИЯ ---
    # Разбиваем текст на слова и считаем их количество
    words = article.content.split() if article.content else []
    word_count = len(words)
    
    # Средняя скорость чтения — 160 слов в минуту.
    # Округляем и ставим минимум 1 минуту, чтобы не было "0 минут"
    read_time = max(1, round(word_count / 160))
    # -----------------------------

    comment_form = CommentForm()

    bookmarks_count = article.bookmarks.count()
    is_bookmarked = False

    # Ищем статьи из той же категории (максимум 3 штуки), исключая текущую
    similar_articles = Article.objects.filter(category=article.category, is_published=True).exclude(pk=article.pk)[:3]

    if request.user.is_authenticated:
        is_bookmarked = Bookmark.objects.filter(
            user=request.user,
            article=article
        ).exists()

    if request.method == "POST":
        form_type = request.POST.get("form_type")

        if form_type == "rating":
            if not request.user.is_authenticated:
                messages.error(request, "Чтобы поставить оценку, нужно войти в аккаут")
                return redirect("articles:article_detail", pk=article.pk)
            
            if article.author_id == request.user.id:
                messages.error(request, "Автор не может оставить оценку для своей статьи")
                return redirect("articles:article_detail", pk=article.pk)
                
            if ratings.filter(user=request.user).exists():
                messages.error(request, "Вы уже добавили оценку для этой статьи")
                return redirect("articles:article_detail", pk=article.pk)
            
            rating = ArticleRating(
                article=article,
                user = request.user,
                rating = request.POST.get("rating"),
                comment=request.POST.get("comment"),
            )
        
            try:
                rating.full_clean()
                rating.save()

                messages.success(request, "Оценка сохранена")
            except ValidationError as error:
                if hasattr(error, "message_dict"):
                    for field_errors in error.message_dict.values():
                        for text in field_errors:
                            messages.error(request, text)
                else:
                    for text in error.messages:
                        messages.error(request, text)

            return redirect("articles:article_detail", pk=pk)
        
        if form_type == "comment":
            if not request.user.is_authenticated:
                messages.error(request, "Чтобы оставить комментарий, нужно войти в аккаут")
                return redirect("articles:article_detail", pk=article.pk)
            
            comment_form = CommentForm(request.POST)

            if comment_form.is_valid():
                comment = comment_form.save(commit=False)
                comment.article = article
                comment.author_name = request.user.username
                comment.author_email = request.user.email
                comment.is_approved = True
                comment.save()

                if article.author != request.user:
                    Notification.objects.create(
                        recipient=article.author,
                        sender=request.user,
                        notification_type="comment",
                        article=article,
                        comment=comment,
                        message=f"{request.user.username} оставил комментарий к вашей статье"
                    )

                messages.success(
                    request,
                    "Комментарий добавлен"
                )

                return redirect("articles:article_detail", pk=article.pk)

    avg_rating = ratings.aggregate(avg=Avg("rating"))["avg"]
    user_rating = None
    has_rated = False
    is_subscribed = False

    if request.user.is_authenticated:
        user_rating = ratings.filter(user=request.user).first()
        has_rated = user_rating is not None
        # Проверяем подписку на автора
        is_subscribed = Subscription.objects.filter(user=request.user, author=article.author).exists()

    comments = Comment.objects.filter(
        article = article,
        is_approved = True
    )

    context = {
        "article":article,
        "avg_rating":avg_rating,
        "user_rating":user_rating,
        "has_rated":has_rated,
        "rating_choices":ArticleRating.RATING_CHOICES,
        "comment_form": comment_form,
        "comments": comments,
        "bookmarks_count": bookmarks_count,
        "is_bookmarked": is_bookmarked,
        "is_subscribed": is_subscribed,
        "similar_articles": similar_articles,
        "read_time": read_time  # <-- Передали посчитанные минуты в контекст
    }

    return render(request, "articles/article_detail.html", context)

@login_required(login_url="articles:login")
def notification_list(request):
    """Уведомления пользователя"""
    notifications = Notification.objects.filter(
        recipient = request.user
    ).select_related(
        "sender",
        "article",
        "comment"
    ).order_by("is_read", "-created_at")

    return render(
        request,
        "articles/notification.html",
        {
            "notifications" : notifications
        }
    )


@login_required(login_url="articles:login")
@require_POST
def mark_notification_as_read(request, pk):
    notification = get_object_or_404(
        Notification,
        pk=pk,
        recipient = request.user
    )

    notification.mark_as_read()

    return redirect("articles:notification_list")


@login_required(login_url="articles:login")
@require_POST
def mark_all_notification_as_read(request):
    Notification.objects.filter(
        recipient=request.user,
        is_read=False
    ).update(is_read=True)

    return redirect("articles:notification_list")


@require_POST
def toggle_bookmark(request, pk):
    """
    Вернет Json ответ, не реняет HTML страничку
    """
    if not request.user.is_authenticated:
        return JsonResponse(
            {
                "success": False,
                "error": "Чтобы добавить в избранное, нужно авторизироваться"
            },
            status=401
        )
    
    article = get_object_or_404(Article, pk=pk)
    bookmark, created = Bookmark.objects.get_or_create(
        user=request.user,
        article=article
    )

    if created:
        is_bookmarked = True
        message = "Статья добавлена в закладки"
    else:
        bookmark.delete()
        is_bookmarked = False
        message = "Статья удалена из закладок"

    bookmarks_count = article.bookmarks.count()
    
    return JsonResponse(
        {
            "success": True,
            "is_bookmarked": is_bookmarked,
            "bookmarks_count": bookmarks_count,
            "message": message
        }
    )


@login_required(login_url="articles:login")
def my_bookmarks(request):
    bookmarks = Bookmark.objects.filter(
        user=request.user
    ).select_related(
        "article",
        "article__author",
        "article__category"
    ).order_by("-created_at")

    return render(
        request,
        "articles/my_bookmarks.html",
        {
            "bookmarks": bookmarks
        }
    )


@login_required(login_url="articles:login")
def create_article(request):
    if request.method == "POST":
        form = ArticleForm(request.POST)

        if (form.is_valid()):
            article = form.save(commit=False)
            article.author = request.user
            article.save()

            # Отправка уведомлений подписчикам автора
            followers = Subscription.objects.filter(author=request.user)
            for follower in followers:
                Notification.objects.create(
                    recipient=follower.user,
                    sender=request.user,
                    notification_type="comment",  # Используем существующий тип, чтобы не ломать choices
                    article=article,
                    message=f"Автор {request.user.username} опубликовал новую статью: '{article.title}'!"
                )

            return redirect(
                "articles:article_detail",
                pk=article.pk
            )
    else:
        form = ArticleForm()

    return render(
        request,
        "articles/article_create.html",
        {"form": form}
    )


def signup(request):
    if request.method == "POST":
        form = SignUpForm(request.POST)

        if (form.is_valid()):
            user = form.save()
            login(request, user)

            return redirect("articles:article_list")
    else:
        form = SignUpForm()

    return render(
        request,
        "articles/signup.html",
        {"form": form}
    )


@login_required(login_url="articles:login")
def export_articles_to_excel(request):
    """Создает Excel файл со всеми статьями и отправляет пользователю"""

    # 1. Создаем Workbook(книгу) и берем активный лист
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Статьи"

    #2. Первая строка - понятный заголовок
    headers =[
        "ID",
        "Заголовок",
        "Автор",
        "Категория",
        "Дата создания",
        "Дата редактирования",
        "Опубликована"
    ]
    worksheet.append(headers)

    # 3. Получаем статьи
    articles = Article.objects.all()

    #4. Каждая статья -> строчка в excel
    for article in articles:
        category_name = article.category.name if article.category else "Без категория"

        row = [
            article.pk,
            article.title,
            article.author.username,
            category_name,
            article.create_at.strftime("%d.%m.%Y"),
            article.update_at.strftime("%d.%m.%Y"),
            "Да" if article.is_published else "Нет"
        ]
        worksheet.append(row)

    # 4.1 Настраиваем ширину колонок
    widths = [8, 34, 20, 22, 18, 18, 6]
    for column_number, width in enumerate(widths, start=1):
        column_letter = get_column_letter(column_number)
        worksheet.column_dimensions[column_letter].width = width
    
    # 5. Готовим HTTP-ответ, чтобы браузер понял, что нужно скачать Excel
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="articles.xlsx"'

    # 6. Сохраняем Excel
    workbook.save(response)
    return response


@login_required
def toggle_subscription(request, author_id):
    author = get_object_or_404(User, pk=author_id)
    
    # нельзя подписаться на самого себя
    if author == request.user:
        messages.error(request, "Вы не можете подписаться на самого себя!")
        return redirect(request.META.get('HTTP_REFERER', 'articles:article_list'))

    # Проверяем, есть ли уже подписка
    sub, created = Subscription.objects.get_or_create(user=request.user, author=author)

    if not created:
        # Если подписка уже была — удаляем 
        sub.delete()
        messages.success(request, f"Вы отписались от автора {author.username}.")
    else:
        messages.success(request, f"Вы успешно подписались на автора {author.username}!")

    return redirect(request.META.get('HTTP_REFERER', 'articles:article_list'))

@login_required(login_url="articles:login")
def my_subscriptions(request):
    # Получаем все подписки текущего пользователя и сразу подгружаем данные авторов
    subscriptions = Subscription.objects.filter(user=request.user).select_related("author")
    
    return render(
        request, 
        "articles/my_subscriptions.html", 
        {"subscriptions": subscriptions}
    )